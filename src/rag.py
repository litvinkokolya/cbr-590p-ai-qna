import os
import re

import docx
import openai
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings

from src.config import Settings

# Паттерны, по которым определяем начало нового раздела:
# - "Глава 1.", "Глава 12." — заголовки глав
# - "1.1.", "2.3.4." — нумерованные пункты/подпункты
_SECTION_RE = re.compile(r"^(Глава\s+\d+\.|(\d+\.){2,})\s*\S")


def _load_sections(docx_path: str) -> list[Document]:
    """Читает .docx и разбивает на секции по заголовкам глав и нумерованным пунктам.
    Каждый пункт (1.1., 2.3. и т.д.) вместе с дочерними абзацами — один Document."""
    doc = docx.Document(docx_path)

    print('doc paragraphs: ', doc.paragraphs)

    sections: list[Document] = []
    current_header = ""
    current_lines: list[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            print('if text: ', text)
            print('if text current header: ', current_header)
            sections.append(Document(page_content=text, metadata={"header": current_header}))

    for para in doc.paragraphs:
        print('para: ', para)
        text = para.text.strip()
        if not text:
            continue

        if _SECTION_RE.match(text):
            flush()
            current_header = text
            current_lines = [text]
        else:
            current_lines.append(text)

    flush()
    return sections


def build_index(settings: Settings) -> None:
    """Загружает 590-П, нарезает по структуре документа и сохраняет FAISS индекс."""
    print("Загружаю документ и нарезаю по структуре (главы / пункты)...")
    chunks = _load_sections(settings.docx_path)
    print(f"Секций: {len(chunks)}")

    print("Создаю эмбеддинги и индекс (запрос к OpenAI)...")
    embeddings = OpenAIEmbeddings(model=settings.embeddings_model, api_key=settings.openai_api_key)
    try:
        vectorstore = FAISS.from_documents(chunks, embeddings)
    except openai.AuthenticationError:
        raise RuntimeError("Ошибка аутентификации OpenAI: проверьте API ключ.")
    except openai.RateLimitError:
        raise RuntimeError("Превышен лимит запросов OpenAI при создании индекса. Попробуйте позже.")
    except openai.APIConnectionError:
        raise RuntimeError("Не удалось подключиться к OpenAI для создания индекса.")
    vectorstore.save_local(settings.index_path)
    print(f"Индекс сохранён в '{settings.index_path}'")


def get_retriever(settings: Settings):
    """Возвращает retriever — строит индекс если его нет."""
    if not os.path.exists(settings.index_path):
        build_index(settings)

    embeddings = OpenAIEmbeddings(model=settings.embeddings_model, api_key=settings.openai_api_key)
    vectorstore = FAISS.load_local(
        settings.index_path,
        embeddings,
        allow_dangerous_deserialization=True,
    )
    return vectorstore.as_retriever(search_kwargs={"k": settings.retriever_k})
